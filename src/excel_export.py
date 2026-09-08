"""Exportació de l'informe mensual a Excel (només informe, no BD)."""
from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _autosize(ws) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[letter].width = min(width + 2, 40)
    for cell in ws[1]:
        cell.font = Font(bold=True)


def generar_excel(df_mes: pd.DataFrame, resum: dict, any_mes: str) -> BytesIO:
    buf = BytesIO()
    movs = df_mes.copy()
    if not movs.empty:
        movs["data"] = pd.to_datetime(movs["data"]).dt.strftime("%Y-%m-%d")
        movs = movs[["data", "tipus", "import", "categoria", "subcategoria", "qui", "descripcio", "metode"]]
    else:
        movs = pd.DataFrame(
            columns=["data", "tipus", "import", "categoria", "subcategoria", "qui", "descripcio", "metode"]
        )
    resum_df = pd.DataFrame(
        [
            {"Concepte": "Ingressos", "Import": resum["ingressos"]},
            {"Concepte": "Despeses", "Import": resum["despeses"]},
            {"Concepte": "Estalvi net", "Import": resum["estalvi"]},
            {"Concepte": "Taxa estalvi %", "Import": resum["taxa_estalvi"]},
        ]
    )
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        resum_df.to_excel(writer, sheet_name="Resum", index=False)
        resum["per_categoria"].to_excel(writer, sheet_name="Per Categoria", index=False)
        resum["per_persona"].to_excel(writer, sheet_name="Per Persona", index=False)
        movs.to_excel(writer, sheet_name="Moviments", index=False)
        for ws in writer.sheets.values():
            _autosize(ws)
        writer.sheets["Resum"]["A1"] = f"Resum {any_mes}"
    buf.seek(0)
    return buf
